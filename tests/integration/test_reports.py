"""Báo cáo: số liệu khớp DB, xuất HTML/CSV/Excel an toàn (escape, chống formula injection, không lộ secret),
phân quyền. Lần chạy được thực thi thật bằng worker executor với demo-target."""

import csv
import io
import uuid
from datetime import date, timedelta
from typing import Any

import httpx
from omnitest_worker.executor import RunExecutor
from openpyxl import load_workbook
from sqlalchemy import func, select

from omnitest_api.db import get_engine, get_sessionmaker
from omnitest_api.main import create_app
from omnitest_api.models import AuditLog, TestRun
from tests.conftest import LoginAs, MakeUser
from tests.integration.phase3_helpers import API_HEALTH, API_LOGIN_FAIL, case_body, setup_project

SECRET = "Demo@123"
XSS_TITLE = "<script>alert('x')</script> & \"trích\""
FORMULA_TITLE = '=HYPERLINK("http://evil.example","bấm")'
FAILING: dict[str, Any] = {
    "kind": "api",
    "steps": [
        {
            "name": "Không token",
            "url": "/api/items?mat_khau={{DEMO_PASSWORD}}",
            "assertions": [{"type": "status", "equals": 200}],
        }
    ],
}


async def _case(c: httpx.AsyncClient, pid: str, title: str, definition: dict[str, Any]) -> str:
    r = await c.post(f"/api/projects/{pid}/test-cases", json=case_body(title, definition))
    assert r.status_code == 201, r.text
    return str(r.json()["id"])


async def _run(
    c: httpx.AsyncClient, pid: str, eid: str, ids: list[str] | None = None, suite: str | None = None
) -> str:
    body: dict[str, Any] = {"environment_id": eid, "confirm": True}
    if suite:
        body["suite_id"] = suite
    else:
        body["test_case_ids"] = ids
    r = await c.post(f"/api/projects/{pid}/runs", json=body)
    assert r.status_code == 201, r.text
    run_id = str(r.json()["id"])
    await RunExecutor(get_engine()).execute_run(uuid.UUID(run_id))
    return run_id


async def test_run_report_formats_are_correct_and_safe(
    make_user: MakeUser, login_as: LoginAs, demo_url: str
) -> None:
    await make_user("root@x.vn", admin=True)
    root = await login_as("root@x.vn")
    pid, eid = await setup_project(root, demo_url)
    ok = await _case(root, pid, XSS_TITLE, API_HEALTH)
    bad = await _case(root, pid, FORMULA_TITLE, FAILING)
    login_fail = await _case(root, pid, "Sai mật khẩu", API_LOGIN_FAIL)
    run_id = await _run(root, pid, eid, [ok, bad, login_fail])
    detail = (await root.get(f"/api/runs/{run_id}")).json()
    statuses = [i["status"] for i in detail["items"]]

    r = await root.get(f"/api/runs/{run_id}/report")
    assert r.status_code == 200
    assert r.headers["content-type"].startswith("text/html")
    assert "sandbox" in r.headers["content-security-policy"]
    page = r.text
    assert "<script>" not in page and "&lt;script&gt;alert(&#x27;x&#x27;)&lt;/script&gt;" in page
    assert SECRET not in page
    assert f"Đạt: {detail['passed']}" in page and f"Không đạt: {detail['failed']}" in page

    r = await root.get(f"/api/runs/{run_id}/report", params={"format": "csv"})
    assert r.status_code == 200 and "attachment" in r.headers["content-disposition"]
    text = r.content.decode("utf-8-sig")
    assert SECRET not in text
    rows = list(csv.reader(io.StringIO(text)))
    assert rows[0][0] == "STT" and len(rows) == 1 + len(detail["items"])
    titles = [row[2] for row in rows[1:]]
    assert titles[0] == XSS_TITLE
    # Ô bắt đầu bằng "=" bị vô hiệu hóa để Excel không hiểu là công thức.
    assert titles[1] == "'" + FORMULA_TITLE
    labels = {"passed": "Đạt", "failed": "Không đạt"}
    assert [row[5] for row in rows[1:]] == [labels[s] for s in statuses]

    r = await root.get(f"/api/runs/{run_id}/report", params={"format": "xlsx"})
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Tổng quan", "Kết quả test case", "Bước không đạt"]
    items = list(wb["Kết quả test case"].iter_rows(min_row=2, values_only=True))
    assert len(items) == 3 and items[1][2] == "'" + FORMULA_TITLE
    steps = list(wb["Bước không đạt"].iter_rows(min_row=2, values_only=True))
    assert {s[0] for s in steps} == {i["code"] for i in detail["items"] if i["status"] != "passed"}
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            assert all(SECRET not in str(v) for v in row if v is not None)

    assert (
        await root.get(f"/api/runs/{run_id}/report", params={"format": "pdf"})
    ).status_code == 422
    async with get_sessionmaker()() as s:
        n = await s.scalar(
            select(func.count()).where(
                AuditLog.action == "report.export", AuditLog.target_id == run_id
            )
        )
    assert n == 3


async def test_project_summary_matches_database_and_filters(
    make_user: MakeUser, login_as: LoginAs, demo_url: str
) -> None:
    await make_user("root@x.vn", admin=True)
    root = await login_as("root@x.vn")
    pid, dev = await setup_project(root, demo_url)
    staging = (
        await root.post(
            f"/api/projects/{pid}/environments",
            json={
                "name": "Staging",
                "kind": "staging",
                "base_url": demo_url,
                "allow_private_network": True,
            },
        )
    ).json()["id"]
    await root.post(
        f"/api/environments/{staging}/variables",
        json={"key": "DEMO_PASSWORD", "value": SECRET, "is_secret": True},
    )
    ok = await _case(root, pid, "Health", API_HEALTH)
    bad = await _case(root, pid, "Không token", FAILING)
    suite = (
        await root.post(
            f"/api/projects/{pid}/suites", json={"name": "Smoke", "test_case_ids": [ok, bad]}
        )
    ).json()["id"]
    await _run(root, pid, dev, suite=suite)
    await _run(root, pid, dev, [ok])
    await _run(root, pid, staging, [bad])

    # Đối chiếu với bộ đếm lưu trong DB của từng lần chạy.
    async with get_sessionmaker()() as s:
        db_runs = list(await s.scalars(select(TestRun).where(TestRun.project_id == uuid.UUID(pid))))
    db_passed = sum(r.passed for r in db_runs)
    db_failed = sum(r.failed for r in db_runs)
    assert len(db_runs) == 3 and (db_passed, db_failed) == (2, 2)

    rep = (await root.get(f"/api/projects/{pid}/reports/summary")).json()
    assert rep["runs"] == 3 and rep["run_statuses"] == {"failed": 2, "passed": 1}
    assert rep["items"]["passed"] == db_passed and rep["items"]["failed"] == db_failed
    assert rep["pass_rate"] == 50.0
    by_env = {g["label"]: g for g in rep["by_environment"]}
    assert by_env["Dev"]["runs"] == 2 and by_env["Dev"]["items"]["passed"] == 2
    assert by_env["Staging"]["items"]["failed"] == 1 and by_env["Staging"]["pass_rate"] == 0.0
    by_suite = {g["label"]: g for g in rep["by_suite"]}
    assert by_suite["Smoke"]["runs"] == 1 and by_suite["Smoke"]["pass_rate"] == 50.0
    assert sum(g["runs"] for g in rep["by_suite"]) == 3
    cases = {c["title"]: c for c in rep["test_cases"]}
    assert cases["Health"]["executions"] == 2 and cases["Health"]["pass_rate"] == 100.0
    assert cases["Không token"]["items"]["failed"] == 2
    assert rep["test_cases"][0]["title"] == "Không token"  # lỗi nhiều nhất lên đầu
    assert sum(d["items"]["passed"] + d["items"]["failed"] for d in rep["daily"]) == 4
    assert len(rep["daily"]) == 30

    only_staging = (
        await root.get(f"/api/projects/{pid}/reports/summary", params={"environment_id": staging})
    ).json()
    assert only_staging["runs"] == 1 and only_staging["filters"]["environment_name"] == "Staging"
    only_suite = (
        await root.get(f"/api/projects/{pid}/reports/summary", params={"suite_id": suite})
    ).json()
    assert only_suite["runs"] == 1 and only_suite["items"]["passed"] == 1

    future = (date.today() + timedelta(days=5)).isoformat()
    empty = (
        await root.get(
            f"/api/projects/{pid}/reports/summary", params={"date_from": future, "date_to": future}
        )
    ).json()
    assert empty["runs"] == 0 and empty["pass_rate"] is None and empty["test_cases"] == []

    bad_range = await root.get(
        f"/api/projects/{pid}/reports/summary",
        params={"date_from": "2026-02-01", "date_to": "2026-01-01"},
    )
    assert bad_range.status_code == 422
    too_long = await root.get(
        f"/api/projects/{pid}/reports/summary",
        params={"date_from": "2024-01-01", "date_to": "2026-01-01"},
    )
    assert too_long.status_code == 422

    xlsx = await root.get(f"/api/projects/{pid}/reports/summary/export")
    wb = load_workbook(io.BytesIO(xlsx.content))
    assert wb.sheetnames == [
        "Tổng quan",
        "Theo môi trường",
        "Theo suite",
        "Theo test case",
        "Theo ngày",
    ]
    env_rows = {r[0]: r for r in wb["Theo môi trường"].iter_rows(min_row=2, values_only=True)}
    assert env_rows["Dev"][1] == 2 and env_rows["Staging"][3] == 1
    csv_text = (
        await root.get(f"/api/projects/{pid}/reports/summary/export", params={"format": "csv"})
    ).content.decode("utf-8-sig")
    assert len(list(csv.reader(io.StringIO(csv_text)))) == 3
    html = await root.get(f"/api/projects/{pid}/reports/summary/export", params={"format": "html"})
    assert "Theo môi trường" in html.text and SECRET not in html.text


async def test_report_access_control(make_user: MakeUser, login_as: LoginAs, demo_url: str) -> None:
    await make_user("root@x.vn", admin=True)
    await make_user("viewer@x.vn")
    await make_user("out@x.vn")
    root = await login_as("root@x.vn")
    pid, eid = await setup_project(root, demo_url)
    other_pid, other_eid = await setup_project(root, demo_url, key="KHAC")
    r = await root.post(
        f"/api/projects/{pid}/members", json={"email": "viewer@x.vn", "role": "viewer"}
    )
    assert r.status_code == 201, r.text
    run_id = await _run(root, pid, eid, [await _case(root, pid, "Health", API_HEALTH)])

    viewer = await login_as("viewer@x.vn")
    assert (await viewer.get(f"/api/runs/{run_id}/report")).status_code == 200
    assert (await viewer.get(f"/api/projects/{pid}/reports/summary")).status_code == 200

    outsider = await login_as("out@x.vn")
    for url in (
        f"/api/runs/{run_id}/report",
        f"/api/projects/{pid}/reports/summary",
        f"/api/projects/{pid}/reports/summary/export",
    ):
        assert (await outsider.get(url)).status_code == 404
    # Bộ lọc trỏ tới môi trường của dự án khác (IDOR) bị từ chối, không lộ tên môi trường.
    r = await root.get(f"/api/projects/{pid}/reports/summary", params={"environment_id": other_eid})
    assert r.status_code == 404
    assert (await root.get(f"/api/projects/{other_pid}/reports/summary")).json()["runs"] == 0

    async with httpx.AsyncClient(
        transport=httpx.ASGITransport(app=create_app()), base_url="http://test"
    ) as anon:
        assert (await anon.get(f"/api/runs/{run_id}/report")).status_code == 401
