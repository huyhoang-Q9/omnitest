import io

from openpyxl import load_workbook

from tests.conftest import LoginAs, MakeUser
from tests.integration.phase3_helpers import (
    API_HEALTH,
    API_LOGIN_FAIL,
    WEB_LOGIN,
    case_body,
    setup_project,
)


async def _root(make_user: MakeUser, login_as: LoginAs):  # type: ignore[no-untyped-def]
    await make_user("root@x.vn", admin=True)
    return await login_as("root@x.vn")


async def test_create_codes_versions_and_side_effect(
    make_user: MakeUser, login_as: LoginAs
) -> None:
    root = await _root(make_user, login_as)
    pid, _ = await setup_project(root, "https://demo.vn")
    a = (
        await root.post(
            f"/api/projects/{pid}/test-cases",
            json=case_body("Health", API_HEALTH, tags=["Smoke", "smoke"]),
        )
    ).json()
    b = (
        await root.post(
            f"/api/projects/{pid}/test-cases", json=case_body("Login sai", API_LOGIN_FAIL)
        )
    ).json()
    assert a["code"] == "DEMO-TC-0001" and b["code"] == "DEMO-TC-0002"
    assert a["tags"] == ["smoke"] and a["side_effect"] is False and a["step_count"] == 1
    # POST được tự đánh dấu side effect.
    assert b["side_effect"] is True

    r = await root.patch(
        f"/api/test-cases/{a['id']}",
        json={"version": 1, "title": "Kiểm tra health", "change_note": "đổi tên"},
    )
    assert r.status_code == 200 and r.json()["version"] == 2
    # Sửa với version cũ bị từ chối (khóa lạc quan).
    stale = await root.patch(f"/api/test-cases/{a['id']}", json={"version": 1, "title": "x"})
    assert stale.status_code == 409
    # Không đổi gì thì không tăng version.
    same = await root.patch(
        f"/api/test-cases/{a['id']}", json={"version": 2, "title": "Kiểm tra health"}
    )
    assert same.json()["version"] == 2
    versions = (await root.get(f"/api/test-cases/{a['id']}/versions")).json()
    assert [v["version"] for v in versions] == [2, 1]
    assert versions[0]["changed_fields"] == ["title"] and versions[0]["change_note"] == "đổi tên"


async def test_definition_validation(make_user: MakeUser, login_as: LoginAs) -> None:
    root = await _root(make_user, login_as)
    pid, _ = await setup_project(root, "https://demo.vn")
    wrong_kind = await root.post(
        f"/api/projects/{pid}/test-cases",
        json={"title": "x", "test_type": "api", "definition": WEB_LOGIN},
    )
    assert wrong_kind.status_code == 422
    sleep = {"kind": "web", "steps": [{"name": "Ngủ", "action": "sleep", "value": "5"}]}
    assert (
        await root.post(f"/api/projects/{pid}/test-cases", json=case_body("x", sleep))
    ).status_code == 422
    desktop = await root.post(
        f"/api/projects/{pid}/test-cases",
        json=case_body("Mac", {"kind": "desktop_macos", "notes": "Mở app"}),
    )
    assert desktop.status_code == 201


async def test_duplicate_delete_restore_and_filters(make_user: MakeUser, login_as: LoginAs) -> None:
    root = await _root(make_user, login_as)
    pid, _ = await setup_project(root, "https://demo.vn")
    a = (
        await root.post(
            f"/api/projects/{pid}/test-cases",
            json=case_body("Health", API_HEALTH, priority="high", tags=["smoke"]),
        )
    ).json()
    await root.post(f"/api/projects/{pid}/test-cases", json=case_body("Web login", WEB_LOGIN))
    dup = (await root.post(f"/api/test-cases/{a['id']}/duplicate")).json()
    assert (
        dup["code"] == "DEMO-TC-0003"
        and dup["title"].endswith("(bản sao)")
        and dup["status"] == "draft"
    )

    def q(**params: object):  # type: ignore[no-untyped-def]
        return root.get(f"/api/projects/{pid}/test-cases", params=params)

    assert (await q(test_type="web")).json()["total"] == 1
    assert (await q(priority="high")).json()["total"] == 2
    assert (await q(tag="smoke")).json()["total"] == 2
    assert (await q(q="tc-0002")).json()["total"] == 1

    assert (await root.delete(f"/api/test-cases/{a['id']}")).status_code == 204
    assert (await q()).json()["total"] == 2
    assert (await q(include_deleted=True)).json()["total"] == 3
    assert (
        await root.patch(f"/api/test-cases/{a['id']}", json={"version": 1, "title": "x"})
    ).status_code == 409
    assert (await root.post(f"/api/test-cases/{a['id']}/restore")).json()["is_deleted"] is False


async def test_viewer_and_tester_cannot_edit(make_user: MakeUser, login_as: LoginAs) -> None:
    root = await _root(make_user, login_as)
    pid, _ = await setup_project(root, "https://demo.vn")
    await make_user("t@x.vn")
    await root.post(f"/api/projects/{pid}/members", json={"email": "t@x.vn", "role": "tester"})
    t = await login_as("t@x.vn")
    assert (
        await t.post(f"/api/projects/{pid}/test-cases", json=case_body("x", API_HEALTH))
    ).status_code == 403
    assert (await t.get(f"/api/projects/{pid}/test-cases")).status_code == 200


async def test_export_import_roundtrip_csv_and_xlsx(make_user: MakeUser, login_as: LoginAs) -> None:
    root = await _root(make_user, login_as)
    pid, _ = await setup_project(root, "https://demo.vn")
    await root.post(
        f"/api/projects/{pid}/test-cases", json=case_body('=HYPERLINK("x")', API_HEALTH)
    )
    await root.post(f"/api/projects/{pid}/test-cases", json=case_body("Web", WEB_LOGIN))

    csv_resp = await root.get(f"/api/projects/{pid}/test-cases/export", params={"format": "csv"})
    assert csv_resp.status_code == 200 and "attachment" in csv_resp.headers["content-disposition"]
    text = csv_resp.content.decode("utf-8-sig")
    # Chống formula injection: ô bắt đầu bằng = được thêm dấu '.
    assert "'=HYPERLINK" in text

    xlsx = await root.get(f"/api/projects/{pid}/test-cases/export", params={"format": "xlsx"})
    ws = load_workbook(io.BytesIO(xlsx.content)).active
    assert ws is not None and ws.max_row == 3

    # Import vào dự án khác: dry run trước, rồi ghi thật. Bỏ cột code để tạo mới.
    pid2 = (await root.post("/api/projects", json={"key": "OTHER", "name": "O"})).json()["id"]
    lines = text.splitlines()
    header = lines[0].split(",")
    assert header[0] == "code"
    no_code = "\n".join([lines[0], *[("" + line[line.index(",") :]) for line in lines[1:]]])
    files = {"file": ("cases.csv", no_code.encode("utf-8"), "text/csv")}
    dry = (await root.post(f"/api/projects/{pid2}/test-cases/import", files=files)).json()
    assert dry == {"dry_run": True, "total_rows": 2, "created": 2, "updated": 0, "errors": []}
    assert (await root.get(f"/api/projects/{pid2}/test-cases")).json()["total"] == 0
    real = (
        await root.post(
            f"/api/projects/{pid2}/test-cases/import", params={"dry_run": False}, files=files
        )
    ).json()
    assert real["dry_run"] is False and real["created"] == 2
    items = (await root.get(f"/api/projects/{pid2}/test-cases")).json()["items"]
    assert items[0]["title"] == '=HYPERLINK("x")' and items[0]["code"] == "OTHER-TC-0001"

    # Import xlsx cập nhật theo mã có sẵn.
    files2 = {"file": ("cases.xlsx", xlsx.content, "application/octet-stream")}
    upd = (
        await root.post(
            f"/api/projects/{pid}/test-cases/import", params={"dry_run": False}, files=files2
        )
    ).json()
    assert upd["updated"] == 2 and upd["created"] == 0 and upd["errors"] == []


async def test_import_reports_row_errors_and_writes_nothing(
    make_user: MakeUser, login_as: LoginAs
) -> None:
    root = await _root(make_user, login_as)
    pid, _ = await setup_project(root, "https://demo.vn")
    csv_text = (
        "title,test_type,definition_json\n"
        'Tốt,api,"{""kind"":""api"",""steps"":[{""name"":""h"",""url"":""/""}]}"\n'
        "Hỏng,api,khong-phai-json\n"
        'Sai loại,web,"{""kind"":""api"",""steps"":[{""name"":""h"",""url"":""/""}]}"\n'
    )
    r = (
        await root.post(
            f"/api/projects/{pid}/test-cases/import",
            params={"dry_run": False},
            files={"file": ("x.csv", csv_text.encode(), "text/csv")},
        )
    ).json()
    assert [e["row"] for e in r["errors"]] == [3, 4]
    assert r["dry_run"] is True
    assert (await root.get(f"/api/projects/{pid}/test-cases")).json()["total"] == 0


async def test_import_rejects_bad_files(make_user: MakeUser, login_as: LoginAs) -> None:
    root = await _root(make_user, login_as)
    pid, _ = await setup_project(root, "https://demo.vn")
    url = f"/api/projects/{pid}/test-cases/import"
    assert (
        await root.post(url, files={"file": ("x.exe", b"MZ", "application/octet-stream")})
    ).status_code == 422
    assert (
        await root.post(url, files={"file": ("x.xlsx", b"not a zip", "application/octet-stream")})
    ).status_code == 422
    big = b"title,test_type,definition_json\n" + b"a" * (2 * 1024 * 1024 + 10)
    assert (await root.post(url, files={"file": ("x.csv", big, "text/csv")})).status_code == 413
    assert (
        await root.post(url, files={"file": ("x.csv", b"a,b\n1,2\n", "text/csv")})
    ).status_code == 422


async def test_filters_treat_wildcards_literally_and_unicode_tags(
    make_user: MakeUser, login_as: LoginAs
) -> None:
    root = await _root(make_user, login_as)
    pid, _ = await setup_project(root, "https://demo.vn")
    await root.post(
        f"/api/projects/{pid}/test-cases",
        json=case_body("Giảm 50% giá", API_HEALTH, tags=["thanh-toán"]),
    )
    await root.post(
        f"/api/projects/{pid}/test-cases", json=case_body("Khác", API_HEALTH, tags=["x"])
    )

    def total(**params: object) -> object:
        return root.get(f"/api/projects/{pid}/test-cases", params=params)

    assert (await total(q="50%")).json()["total"] == 1
    assert (await total(q="%")).json()["total"] == 1
    assert (await total(q="_")).json()["total"] == 0
    assert (await total(tag="Thanh-Toán")).json()["total"] == 1
    assert (await total(tag="thanh")).json()["total"] == 0
