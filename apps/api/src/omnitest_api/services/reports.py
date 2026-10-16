"""Báo cáo lần chạy và báo cáo tổng hợp dự án: truy vấn số liệu từ DB và xuất HTML, CSV, Excel.

Nguyên tắc:
- Số liệu lấy trực tiếp từ kết quả từng test case (test_run_items), không ước lượng.
- Tỷ lệ đạt = đạt / (đạt + không đạt + lỗi + bị chặn); None khi chưa có kết quả (không hiển thị 0% giả).
- Mọi chuỗi đưa vào HTML đều được escape; ô CSV/Excel bắt đầu bằng ký tự công thức được thêm dấu ' phía trước.
- Thông báo lỗi trong DB đã được worker che secret trước khi lưu; báo cáo không đọc biến môi trường hay secret.
"""

import csv
import html
import io
import uuid
from collections.abc import Iterable, Sequence
from dataclasses import dataclass
from datetime import UTC, date, datetime, time, timedelta
from typing import Any
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from sqlalchemy import ColumnElement, Select, func, select
from sqlalchemy.dialects.postgresql import distinct_on
from sqlalchemy.ext.asyncio import AsyncSession

from omnitest_api.models import (
    FINAL_RUN_STATUSES,
    Project,
    RunStatus,
    StepResult,
    TestRun,
    TestRunItem,
    TestSuite,
)
from omnitest_api.schemas_reports import (
    DailyStat,
    GroupStat,
    ProjectReportOut,
    ReportFilters,
    StatusCounts,
    TestCaseStat,
)
from omnitest_api.services.test_case_io import safe_cell

STATUS_ORDER = (
    RunStatus.PASSED,
    RunStatus.FAILED,
    RunStatus.ERROR,
    RunStatus.BLOCKED,
    RunStatus.SKIPPED,
    RunStatus.CANCELLED,
)
STATUS_LABELS = {
    RunStatus.QUEUED: "Đang chờ",
    RunStatus.RUNNING: "Đang chạy",
    RunStatus.PASSED: "Đạt",
    RunStatus.FAILED: "Không đạt",
    RunStatus.ERROR: "Lỗi hệ thống",
    RunStatus.BLOCKED: "Bị chặn",
    RunStatus.SKIPPED: "Bỏ qua",
    RunStatus.CANCELLED: "Đã hủy",
}
MAX_RANGE_DAYS = 366
NO_SUITE = "(Chạy lẻ, không theo suite)"


def pass_rate(c: StatusCounts) -> float | None:
    return round(c.passed / c.executed * 100, 1) if c.executed else None


def _count_cols() -> list[ColumnElement[int]]:
    return [func.count().filter(TestRunItem.status == s) for s in STATUS_ORDER]


def _counts(values: Sequence[Any]) -> StatusCounts:
    return StatusCounts(**{s.value: int(v or 0) for s, v in zip(STATUS_ORDER, values, strict=True)})


def _fmt_dt(value: datetime | None, tz: str) -> str:
    return value.astimezone(ZoneInfo(tz)).strftime("%d/%m/%Y %H:%M:%S") if value else ""


def _fmt_ms(ms: int | None) -> str:
    if ms is None:
        return ""
    return f"{ms} ms" if ms < 1000 else f"{ms / 1000:.1f} s"


def _rate(value: float | None) -> str:
    return "Chưa có dữ liệu" if value is None else f"{value:.1f}%".replace(".", ",")


# ---------------------------------------------------------------- Báo cáo tổng hợp dự án


@dataclass
class ProjectReportQuery:
    date_from: date
    date_to: date
    environment_id: uuid.UUID | None = None
    suite_id: uuid.UUID | None = None


def default_range(tz: str, days: int = 30) -> tuple[date, date]:
    today = datetime.now(ZoneInfo(tz)).date()
    return today - timedelta(days=days - 1), today


async def project_report(
    db: AsyncSession,
    project: Project,
    q: ProjectReportQuery,
    tz: str,
    *,
    environment_name: str | None = None,
    suite_name: str | None = None,
) -> ProjectReportOut:
    since = datetime.combine(q.date_from, time.min, ZoneInfo(tz)).astimezone(UTC)
    until = datetime.combine(q.date_to + timedelta(days=1), time.min, ZoneInfo(tz)).astimezone(UTC)
    conds: list[ColumnElement[bool]] = [
        TestRun.project_id == project.id,
        TestRun.status.in_(FINAL_RUN_STATUSES),
        TestRun.finished_at >= since,
        TestRun.finished_at < until,
    ]
    if q.environment_id:
        conds.append(TestRun.environment_id == q.environment_id)
    if q.suite_id:
        conds.append(TestRun.suite_id == q.suite_id)

    def items_query(*cols: Any) -> Select[Any]:
        return select(*cols).join(TestRun, TestRun.id == TestRunItem.run_id).where(*conds)

    run_statuses = {
        s.value: int(n)
        for s, n in (
            await db.execute(
                select(TestRun.status, func.count()).where(*conds).group_by(TestRun.status)
            )
        ).all()
    }
    totals_row = (await db.execute(items_query(*_count_cols()))).one()
    totals = _counts(totals_row)

    env_rows = (
        await db.execute(
            items_query(
                TestRun.environment_name, func.count(func.distinct(TestRun.id)), *_count_cols()
            )
            .group_by(TestRun.environment_name)
            .order_by(TestRun.environment_name)
        )
    ).all()
    by_env = []
    for name, runs, *counts in env_rows:
        c = _counts(counts)
        by_env.append(GroupStat(key=name, label=name, runs=runs, items=c, pass_rate=pass_rate(c)))

    suite_rows = (
        await db.execute(
            items_query(
                TestRun.suite_id,
                TestSuite.name,
                func.count(func.distinct(TestRun.id)),
                *_count_cols(),
            )
            .outerjoin(TestSuite, TestSuite.id == TestRun.suite_id)
            .group_by(TestRun.suite_id, TestSuite.name)
            .order_by(TestSuite.name.nulls_last())
        )
    ).all()
    by_suite = []
    for sid, sname, runs, *counts in suite_rows:
        c = _counts(counts)
        by_suite.append(
            GroupStat(
                key=str(sid) if sid else None,
                label=sname or NO_SUITE,
                runs=runs,
                items=c,
                pass_rate=pass_rate(c),
            )
        )

    # Theo test case (mã ổn định): số lần chạy, kết quả, và kết quả gần nhất.
    case_rows = (
        await db.execute(
            items_query(TestRunItem.code, func.count(), *_count_cols()).group_by(TestRunItem.code)
        )
    ).all()
    latest = {
        item.code: item
        for item in await db.scalars(
            items_query(TestRunItem)
            .ext(distinct_on(TestRunItem.code))
            .order_by(
                TestRunItem.code,
                TestRunItem.finished_at.desc().nulls_last(),
                TestRun.created_at.desc(),
            )
        )
    }
    cases: list[TestCaseStat] = []
    for code, n, *counts in case_rows:
        c = _counts(counts)
        last = latest[code]
        cases.append(
            TestCaseStat(
                test_case_id=last.test_case_id,
                code=code,
                title=last.title,
                executions=n,
                items=c,
                pass_rate=pass_rate(c),
                last_status=last.status,
                last_error=last.error,
                last_run_id=last.run_id,
                last_finished_at=last.finished_at,
            )
        )
    # Test lỗi nhiều nhất lên đầu, sau đó theo mã.
    cases.sort(key=lambda s: (-(s.items.failed + s.items.error + s.items.blocked), s.code))

    day = func.date(func.timezone(tz, TestRun.finished_at))
    daily_rows = {
        d: _counts(counts)
        for d, *counts in (await db.execute(items_query(day, *_count_cols()).group_by(day))).all()
    }
    daily = []
    d = q.date_from
    while d <= q.date_to:
        daily.append(DailyStat(date=d, items=daily_rows.get(d, StatusCounts())))
        d += timedelta(days=1)

    return ProjectReportOut(
        project_id=project.id,
        project_key=project.key,
        project_name=project.name,
        generated_at=datetime.now(UTC),
        filters=ReportFilters(
            date_from=q.date_from,
            date_to=q.date_to,
            timezone=tz,
            environment_id=q.environment_id,
            environment_name=environment_name,
            suite_id=q.suite_id,
            suite_name=suite_name,
        ),
        runs=sum(run_statuses.values()),
        run_statuses=run_statuses,
        items=totals,
        pass_rate=pass_rate(totals),
        by_environment=by_env,
        by_suite=by_suite,
        test_cases=cases,
        daily=daily,
    )


COUNT_HEADERS = ["Đạt", "Không đạt", "Lỗi hệ thống", "Bị chặn", "Bỏ qua", "Đã hủy"]


def _count_values(c: StatusCounts) -> list[int]:
    return [c.passed, c.failed, c.error, c.blocked, c.skipped, c.cancelled]


def _project_tables(r: ProjectReportOut) -> list[tuple[str, list[str], list[list[Any]]]]:
    tz = r.filters.timezone
    group_headers = ["Nhóm", "Số lần chạy", *COUNT_HEADERS, "Tỷ lệ đạt"]
    return [
        (
            "Theo môi trường",
            group_headers,
            [
                [g.label, g.runs, *_count_values(g.items), _rate(g.pass_rate)]
                for g in r.by_environment
            ],
        ),
        (
            "Theo suite",
            group_headers,
            [[g.label, g.runs, *_count_values(g.items), _rate(g.pass_rate)] for g in r.by_suite],
        ),
        (
            "Theo test case",
            [
                "Mã",
                "Tên test case",
                "Số lần chạy",
                *COUNT_HEADERS,
                "Tỷ lệ đạt",
                "Kết quả gần nhất",
                "Lỗi gần nhất",
                "Thời điểm",
            ],
            [
                [
                    c.code,
                    c.title,
                    c.executions,
                    *_count_values(c.items),
                    _rate(c.pass_rate),
                    STATUS_LABELS[c.last_status],
                    c.last_error or "",
                    _fmt_dt(c.last_finished_at, tz),
                ]
                for c in r.test_cases
            ],
        ),
        (
            "Theo ngày",
            ["Ngày", *COUNT_HEADERS],
            [[d.date.strftime("%d/%m/%Y"), *_count_values(d.items)] for d in r.daily],
        ),
    ]


def _project_overview(r: ProjectReportOut) -> list[tuple[str, str]]:
    f = r.filters
    return [
        ("Dự án", f"{r.project_key} · {r.project_name}"),
        ("Khoảng thời gian", f"{f.date_from:%d/%m/%Y} đến {f.date_to:%d/%m/%Y} ({f.timezone})"),
        ("Môi trường", f.environment_name or "Tất cả"),
        ("Suite", f.suite_name or "Tất cả"),
        ("Số lần chạy đã kết thúc", str(r.runs)),
        ("Số kết quả test case", str(r.items.total)),
        ("Tỷ lệ đạt", _rate(r.pass_rate)),
        ("Thời điểm lập báo cáo", _fmt_dt(r.generated_at, f.timezone)),
    ]


def project_report_csv(r: ProjectReportOut) -> bytes:
    """CSV một bảng: kết quả theo test case (dạng dữ liệu phổ biến nhất khi phân tích tiếp)."""
    _, headers, rows = _project_tables(r)[2]
    return _csv(headers, rows)


def project_report_xlsx(r: ProjectReportOut) -> bytes:
    sheets: list[tuple[str, list[str], list[list[Any]]]] = [
        ("Tổng quan", ["Mục", "Giá trị"], [list(x) for x in _project_overview(r)])
    ]
    sheets += _project_tables(r)
    return _xlsx(sheets)


def project_report_html(r: ProjectReportOut) -> bytes:
    body = [
        f"<h1>Báo cáo kiểm thử dự án {_e(r.project_key)}</h1>",
        _kv_table(_project_overview(r)),
        _status_badges(r.items),
    ]
    for title, headers, rows in _project_tables(r):
        body.append(f"<h2>{_e(title)}</h2>")
        body.append(_table(headers, rows) if rows else "<p class='muted'>Không có dữ liệu.</p>")
    return _page(f"Báo cáo {r.project_key}", "".join(body))


# ---------------------------------------------------------------- Báo cáo một lần chạy


@dataclass
class RunReport:
    run: TestRun
    project: Project
    suite_name: str | None
    items: list[TestRunItem]
    # Các bước không đạt của lần thử cuối, theo run_item_id.
    failed_steps: dict[uuid.UUID, list[StepResult]]
    tz: str


async def load_run_report(db: AsyncSession, run: TestRun, project: Project, tz: str) -> RunReport:
    items = list(
        await db.scalars(
            select(TestRunItem).where(TestRunItem.run_id == run.id).order_by(TestRunItem.position)
        )
    )
    suite_name = None
    if run.suite_id:
        suite = await db.get(TestSuite, run.suite_id)
        suite_name = suite.name if suite else None
    failed_steps: dict[uuid.UUID, list[StepResult]] = {}
    attempts = {i.id: i.attempts for i in items}
    if items:
        steps = await db.scalars(
            select(StepResult)
            .where(
                StepResult.run_item_id.in_(attempts),
                StepResult.status.not_in([RunStatus.PASSED, RunStatus.SKIPPED]),
            )
            .order_by(StepResult.run_item_id, StepResult.attempt, StepResult.index)
        )
        for s in steps:
            if s.attempt == attempts.get(s.run_item_id):
                failed_steps.setdefault(s.run_item_id, []).append(s)
    return RunReport(run, project, suite_name, items, failed_steps, tz)


def _run_counts(run: TestRun) -> StatusCounts:
    return StatusCounts(
        passed=run.passed,
        failed=run.failed,
        error=run.errored,
        blocked=run.blocked,
        skipped=run.skipped,
        cancelled=run.cancelled,
    )


def _run_overview(rr: RunReport) -> list[tuple[str, str]]:
    run = rr.run
    duration = (
        int((run.finished_at - run.started_at).total_seconds() * 1000)
        if run.finished_at and run.started_at
        else None
    )
    return [
        ("Dự án", f"{rr.project.key} · {rr.project.name}"),
        ("Lần chạy", run.title),
        ("Mã lần chạy", str(run.id)),
        ("Trạng thái", STATUS_LABELS[run.status]),
        ("Môi trường", run.environment_name),
        ("Suite", rr.suite_name or NO_SUITE),
        ("Người chạy", run.triggered_by_email or ""),
        ("Tạo lúc", _fmt_dt(run.created_at, rr.tz)),
        ("Bắt đầu", _fmt_dt(run.started_at, rr.tz)),
        ("Kết thúc", _fmt_dt(run.finished_at, rr.tz)),
        ("Thời lượng", _fmt_ms(duration)),
        ("Tổng số test", str(run.total)),
        ("Tỷ lệ đạt", _rate(pass_rate(_run_counts(run)))),
        *([("Lỗi lần chạy", run.error)] if run.error else []),
    ]


ITEM_HEADERS = [
    "STT",
    "Mã",
    "Tên test case",
    "Loại",
    "Phiên bản",
    "Kết quả",
    "Số lần thử",
    "Thời lượng",
    "Lỗi",
]


def _item_rows(rr: RunReport) -> list[list[Any]]:
    return [
        [
            i.position + 1,
            i.code,
            i.title,
            i.test_type.value.upper(),
            i.test_case_version,
            STATUS_LABELS[i.status],
            i.attempts,
            _fmt_ms(i.duration_ms),
            i.error or "",
        ]
        for i in rr.items
    ]


STEP_HEADERS = ["Mã test", "Bước", "Tên bước", "Kết quả", "Thời lượng", "Lỗi"]


def _step_rows(rr: RunReport) -> list[list[Any]]:
    rows = []
    for item in rr.items:
        for s in rr.failed_steps.get(item.id, []):
            rows.append(
                [
                    item.code,
                    s.index + 1,
                    s.name,
                    STATUS_LABELS[s.status],
                    _fmt_ms(s.duration_ms),
                    s.error or "",
                ]
            )
    return rows


def run_report_csv(rr: RunReport) -> bytes:
    return _csv(ITEM_HEADERS, _item_rows(rr))


def run_report_xlsx(rr: RunReport) -> bytes:
    return _xlsx(
        [
            ("Tổng quan", ["Mục", "Giá trị"], [list(x) for x in _run_overview(rr)]),
            ("Kết quả test case", ITEM_HEADERS, _item_rows(rr)),
            ("Bước không đạt", STEP_HEADERS, _step_rows(rr)),
        ]
    )


def run_report_html(rr: RunReport) -> bytes:
    steps = _step_rows(rr)
    body = [
        f"<h1>Báo cáo lần chạy: {_e(rr.run.title)}</h1>",
        _kv_table(_run_overview(rr)),
        _status_badges(_run_counts(rr.run)),
        "<h2>Kết quả từng test case</h2>",
        _table(ITEM_HEADERS, _item_rows(rr))
        if rr.items
        else "<p class='muted'>Không có test case.</p>",
        "<h2>Bước không đạt (lần thử cuối)</h2>",
        _table(STEP_HEADERS, steps)
        if steps
        else "<p class='muted'>Không có bước nào không đạt.</p>",
        "<p class='muted'>Ảnh chụp, trace và log đầy đủ xem trong OmniTest (cần đăng nhập).</p>",
    ]
    return _page(f"Báo cáo lần chạy {rr.run.title}", "".join(body))


# ---------------------------------------------------------------- Kết xuất chung


def _e(value: Any) -> str:
    return html.escape(str(value), quote=True)


def _csv(headers: list[str], rows: Iterable[list[Any]]) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(headers)
    for row in rows:
        w.writerow([safe_cell(v) for v in row])
    return ("﻿" + buf.getvalue()).encode("utf-8")


def _xlsx(sheets: list[tuple[str, list[str], list[list[Any]]]]) -> bytes:
    wb = Workbook()
    first = True
    for title, headers, rows in sheets:
        ws = wb.active if first else wb.create_sheet()
        assert ws is not None
        first = False
        ws.title = title[:31]
        ws.append(headers)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for row in rows:
            ws.append([safe_cell(v) for v in row])
        for idx, header in enumerate(headers, start=1):
            longest = max([len(str(header))] + [len(str(r[idx - 1])) for r in rows[:200]])
            ws.column_dimensions[get_column_letter(idx)].width = min(max(10, longest + 2), 60)
        ws.freeze_panes = "A2"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _kv_table(rows: list[tuple[str, str]]) -> str:
    cells = "".join(f"<tr><th>{_e(k)}</th><td>{_e(v)}</td></tr>" for k, v in rows)
    return f"<table class='kv'>{cells}</table>"


def _table(headers: list[str], rows: list[list[Any]]) -> str:
    head = "".join(f"<th>{_e(h)}</th>" for h in headers)
    body = "".join("<tr>" + "".join(f"<td>{_e(v)}</td>" for v in row) + "</tr>" for row in rows)
    return f"<table><thead><tr>{head}</tr></thead><tbody>{body}</tbody></table>"


def _status_badges(c: StatusCounts) -> str:
    parts = [
        f"<span class='b {s.value}'>{_e(label)}: {n}</span>"
        for s, label, n in zip(STATUS_ORDER, COUNT_HEADERS, _count_values(c), strict=True)
    ]
    return f"<p class='badges'>{''.join(parts)}</p>"


CSS = """
body{font-family:system-ui,-apple-system,"Segoe UI",Roboto,sans-serif;color:#1f2937;margin:24px;line-height:1.45}
h1{font-size:22px;margin:0 0 16px}h2{font-size:17px;margin:28px 0 8px}
table{border-collapse:collapse;width:100%;font-size:13px;margin-bottom:8px}
th,td{border:1px solid #d1d5db;padding:6px 8px;text-align:left;vertical-align:top;word-break:break-word}
thead th{background:#f3f4f6}table.kv{width:auto;min-width:420px}table.kv th{background:#f9fafb;width:200px}
.muted{color:#6b7280;font-size:13px}.badges{display:flex;flex-wrap:wrap;gap:8px}
.b{border-radius:999px;padding:3px 10px;font-size:13px;background:#e5e7eb}
.passed{background:#dcfce7}.failed,.error{background:#fee2e2}.blocked{background:#fef3c7}
@media print{body{margin:0}h2{break-after:avoid}tr{break-inside:avoid}}
"""


def _page(title: str, body: str) -> bytes:
    return (
        "<!doctype html><html lang='vi'><head><meta charset='utf-8'>"
        "<meta name='viewport' content='width=device-width,initial-scale=1'>"
        f"<title>{_e(title)}</title><style>{CSS}</style></head><body>{body}"
        "<p class='muted'>Lập bởi OmniTest Automation Platform.</p></body></html>"
    ).encode()
