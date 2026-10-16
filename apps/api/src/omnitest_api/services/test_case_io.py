"""Import/export test case dạng CSV và Excel.

- Export chống CSV/Excel formula injection: ô bắt đầu bằng = + - @ (và tab, CR) được thêm dấu ' phía trước.
- Import kiểm tra toàn bộ trước, chỉ ghi khi không có lỗi nào (tất cả hoặc không gì cả).
"""

import csv
import io
import json
from typing import Any

from openpyxl import Workbook, load_workbook
from pydantic import ValidationError

from omnitest_api.models import TestCase
from omnitest_api.schemas_tests import TestCaseCreate

COLUMNS = [
    "code",
    "title",
    "description",
    "test_type",
    "priority",
    "status",
    "tags",
    "preconditions",
    "expected_result",
    "side_effect",
    "retry_count",
    "timeout_seconds",
    "test_data_json",
    "definition_json",
]
HEADER_LABELS = {
    "code": "Mã (để trống khi tạo mới)",
}
MAX_ROWS = 2000
DANGEROUS_PREFIX = ("=", "+", "-", "@", "\t", "\r")


class ImportFileError(Exception):
    pass


def safe_cell(value: Any) -> Any:
    if isinstance(value, str) and value.startswith(DANGEROUS_PREFIX):
        return "'" + value
    return value


def _row(tc: TestCase) -> list[Any]:
    return [
        tc.code,
        tc.title,
        tc.description,
        tc.test_type.value,
        tc.priority.value,
        tc.status.value,
        ", ".join(tc.tags),
        tc.preconditions,
        tc.expected_result,
        "true" if tc.side_effect else "false",
        tc.retry_count,
        tc.timeout_seconds,
        json.dumps(tc.test_data, ensure_ascii=False),
        json.dumps(tc.definition, ensure_ascii=False),
    ]


def export_csv(cases: list[TestCase]) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(COLUMNS)
    for tc in cases:
        w.writerow([safe_cell(v) for v in _row(tc)])
    # BOM để Excel mở đúng tiếng Việt.
    return ("﻿" + buf.getvalue()).encode("utf-8")


def export_xlsx(cases: list[TestCase]) -> bytes:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Test cases"
    ws.append(COLUMNS)
    for tc in cases:
        ws.append([safe_cell(v) for v in _row(tc)])
    for col, width in zip(
        "ABCDEFGHIJKLMN", (16, 40, 40, 12, 10, 10, 20, 30, 30, 10, 10, 12, 30, 60), strict=True
    ):
        ws.column_dimensions[col].width = width
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _unescape(value: Any) -> Any:
    if isinstance(value, str) and value.startswith("'") and value[1:2] in DANGEROUS_PREFIX:
        return value[1:]
    return value


def read_rows(filename: str, data: bytes) -> list[dict[str, Any]]:
    name = filename.lower()
    if name.endswith(".csv"):
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise ImportFileError("File CSV phải mã hóa UTF-8") from exc
        reader = csv.DictReader(io.StringIO(text))
        rows = [dict(r) for r in reader]
        header = reader.fieldnames or []
    elif name.endswith(".xlsx"):
        if not data.startswith(b"PK"):
            raise ImportFileError("File không phải định dạng Excel .xlsx hợp lệ")
        try:
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception as exc:
            raise ImportFileError("Không đọc được file Excel") from exc
        ws = wb.worksheets[0]
        values = list(ws.iter_rows(values_only=True))
        if not values:
            raise ImportFileError("File Excel trống")
        header = [str(h).strip() if h is not None else "" for h in values[0]]
        rows = [
            {
                header[i]: (cell if cell is not None else "")
                for i, cell in enumerate(r)
                if i < len(header)
            }
            for r in values[1:]
            if any(c not in (None, "") for c in r)
        ]
        wb.close()
    else:
        raise ImportFileError("Chỉ hỗ trợ file .csv hoặc .xlsx")
    missing = [c for c in ("title", "test_type", "definition_json") if c not in header]
    if missing:
        raise ImportFileError(f"Thiếu cột bắt buộc: {', '.join(missing)}")
    if len(rows) > MAX_ROWS:
        raise ImportFileError(f"Tối đa {MAX_ROWS} dòng mỗi lần import")
    return [{k: _unescape(v) for k, v in r.items() if k} for r in rows]


def parse_row(row: dict[str, Any]) -> tuple[str | None, TestCaseCreate]:
    """Trả (code, dữ liệu). Ném ValueError với thông điệp tiếng Việt khi dòng sai."""

    def text(key: str) -> str:
        v = row.get(key, "")
        return "" if v is None else str(v).strip()

    try:
        definition = json.loads(text("definition_json"))
    except json.JSONDecodeError as exc:
        raise ValueError(f"definition_json không phải JSON hợp lệ ({exc.msg})") from exc
    test_data_raw = text("test_data_json")
    try:
        test_data = json.loads(test_data_raw) if test_data_raw else {}
    except json.JSONDecodeError as exc:
        raise ValueError(f"test_data_json không phải JSON hợp lệ ({exc.msg})") from exc

    payload: dict[str, Any] = {
        "title": text("title"),
        "description": text("description"),
        "test_type": text("test_type"),
        "tags": [t.strip() for t in text("tags").split(",") if t.strip()],
        "preconditions": text("preconditions"),
        "expected_result": text("expected_result"),
        "side_effect": text("side_effect").lower() in {"true", "1", "yes", "x", "có"},
        "test_data": test_data,
        "definition": definition,
    }
    for key in ("priority", "status"):
        if text(key):
            payload[key] = text(key)
    for key in ("retry_count", "timeout_seconds"):
        if text(key):
            try:
                payload[key] = int(float(text(key)))
            except ValueError as exc:
                raise ValueError(f"{key} phải là số") from exc
    try:
        data = TestCaseCreate.model_validate(payload)
    except ValidationError as exc:
        first = exc.errors()[0]
        loc = ".".join(str(x) for x in first["loc"])
        raise ValueError(f"{loc}: {first['msg'].removeprefix('Value error, ')}") from exc
    return (text("code") or None), data
